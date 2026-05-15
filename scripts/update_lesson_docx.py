# -*- coding: utf-8 -*-
"""
更新公开课教案与导学案：
1. 教案：在教学过程表格开头插入"0. 课前进入"行，并相应微调时间。
2. 导学案：末尾追加预印思维导图骨架。
两份文件都直接在 docs/ 下原地更新。
"""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
LESSON_DOCX = ROOT / 'docs' / '信息系统复习课教案.docx'
GUIDE_DOCX = ROOT / 'docs' / '信息系统复习课导学案.docx'


def _set_cell_text(cell, text):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.font.size = Pt(11)


def update_lesson_plan(path: Path) -> None:
    doc = Document(path)

    # 找到教学过程表格（5列：环节/时间/教师/学生/知识点）
    target_table = None
    for table in doc.tables:
        if len(table.columns) == 5 and table.rows[0].cells[0].text.strip() == '环节':
            target_table = table
            break

    if target_table is None:
        raise RuntimeError('未找到教学过程表格')

    # 幂等：如果尚未插入"0. 课前进入"行，则插入
    existing_titles = {row.cells[0].text.strip() for row in target_table.rows}
    if '0. 课前进入' not in existing_titles:
        template_row = target_table.rows[1]._tr
        new_row_xml = deepcopy(template_row)
        template_row.addprevious(new_row_xml)
        new_row = target_table.rows[1]
        new_contents = [
            '0. 课前进入',
            '3分钟',
            '提前 3 分钟到位，启动屏幕广播或教学系统；引导学生打开浏览器进入 /openclass 公开课入口，确认画布加载完成，准备好导学案与笔。',
            '在机房电脑上打开浏览器，访问公开课入口 /openclass；确认画布已经加载基础组件；翻开导学案准备记录。',
            '提示学生公开课入口免登录直接进入，提交时再实名；为后续 5 个环节预留稳定的设备和网络状态。',
        ]
        for cell, text in zip(new_row.cells, new_contents):
            _set_cell_text(cell, text)

    # 调整后续行时间：情境导入 5→4，硬件 10→9，软件 8→7，数据 8→7。其余不变。
    time_updates = {
        '1. 情境导入': '4分钟',
        '2. 硬件搭建': '9分钟',
        '3. 软件分析': '7分钟',
        '5. 数据观察与故障反查': '7分钟',
    }
    for row in target_table.rows[1:]:
        title = row.cells[0].text.strip()
        new_time = time_updates.get(title)
        if new_time:
            _set_cell_text(row.cells[1], new_time)

    # 同步更新四、教学准备中"预设代码陷阱"的表述（去掉揭谜口径）
    target_text = (
        '预设代码陷阱：starter 代码中传感器读取写为 pin0.read_analog()，'
        '蜂鸣器控制写为 pin3.write_digital(...)，代码注释只提示"检查引脚是否一致"，'
        '需要学生自己回到画布看出 P1 / P2 才能修正。'
    )
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith('预设代码陷阱'):
            for run in paragraph.runs:
                run.text = ''
            paragraph.runs[0].text = target_text
            break

    # 同步更新课时一栏：40-45 → 40 分钟
    for table in doc.tables:
        if len(table.rows) >= 2 and table.rows[0].cells[0].text.strip() == '课题':
            _set_cell_text(table.rows[1].cells[2], '1课时（40分钟）')
            break

    doc.save(path)
    print(f'已更新教案：{path}')


def update_guide(path: Path) -> None:
    doc = Document(path)

    # 同步去除导学案中的剧透：段落 + 表格 cell 中的代码注释
    for paragraph in doc.paragraphs:
        if '故意与自动连线不一致' in paragraph.text:
            new_text = (
                '先运行测试，再结合画布检查代码。下面两处代码不一定与画布连线一致，请逐行核对。'
            )
            for run in paragraph.runs:
                run.text = ''
            if paragraph.runs:
                paragraph.runs[0].text = new_text
            else:
                paragraph.add_run(new_text)
            break

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if 'pin0.read_analog' in cell.text and '画布上' in cell.text:
                    _set_cell_text(
                        cell,
                        'raw = pin0.read_analog()       # 这里的引脚号是否和画布上的传感器 DATA 接口一致？\n'
                        'pin3.write_digital(1)         # 这里的引脚号是否和画布上的蜂鸣器 IO 接口一致？'
                    )

    # 幂等清理：删除原"九、思维导图"或之前生成过的"十、"附录，连带后续表格/段落一并清除。
    body = doc.element.body
    children = list(body)
    cleanup_start = None
    for index, element in enumerate(children):
        if element.tag != qn('w:p'):
            continue
        text = ''.join(t.text or '' for t in element.iter(qn('w:t')))
        text = text.strip()
        if text.startswith('九、') or text.startswith('十、'):
            cleanup_start = index
            break
    if cleanup_start is not None:
        sect_pr = body.find(qn('w:sectPr'))
        for element in children[cleanup_start:]:
            if element is sect_pr:
                continue
            body.remove(element)

    # 追加：新的"九、预印思维导图骨架"
    doc.add_heading('九、预印思维导图骨架', level=1)

    intro = doc.add_paragraph()
    intro.add_run(
        '中心节点是"教室温度检测系统"。下面已为你预留了 5 个分支，每个分支已经给出一个'
        '范例叶子，请你在剩下的 2 个空格上补全。完成后再把你认为最容易混淆的 2 个点圈出来。'
    )

    # 思维导图骨架：用 6 列表格表示中心 + 5 分支
    table = doc.add_table(rows=4, cols=6)
    table.style = 'Light Grid Accent 1'

    headers = ['中心', '人', '硬件', '软件', '数据', '网络']
    for cell, text in zip(table.rows[0].cells, headers):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        _set_cell_text(cell, text)

    # 中心列在 3 行内填"教室温度检测系统"（合并）
    center_cells = [row.cells[0] for row in table.rows[1:]]
    merged_center = center_cells[0].merge(center_cells[-1])
    _set_cell_text(merged_center, '教室温度检测系统')
    merged_center.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    sample_leaves = [
        '例：开发者',
        '例：温湿度传感器',
        '例：micro:bit 代码',
        '例：sensorlog 表',
        '例：POST /upload',
    ]
    for cell, text in zip(table.rows[1].cells[1:], sample_leaves):
        _set_cell_text(cell, text)

    for row in table.rows[2:]:
        for cell in row.cells[1:]:
            _set_cell_text(cell, '__________________')

    doc.add_paragraph()
    tip = doc.add_paragraph()
    tip.add_run('补叶提示：').bold = True
    tip.add_run(
        '"人"想想还有谁会用到这个系统？"硬件"还有什么是必须的设备？"软件"除了 micro:bit 端代码还有哪一段？'
        '"数据"除了 sensorlog 还会保存什么？"网络"除了上传还有哪种 HTTP 请求？'
    )

    doc.add_paragraph('我觉得本节课最容易混淆的 2 个知识点：')
    doc.add_paragraph('1. _______________________________________________________________________')
    doc.add_paragraph('2. _______________________________________________________________________')

    doc.save(path)
    print(f'已更新导学案：{path}')


if __name__ == '__main__':
    update_lesson_plan(LESSON_DOCX)
    update_guide(GUIDE_DOCX)

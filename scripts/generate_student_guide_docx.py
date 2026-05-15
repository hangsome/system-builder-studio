from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PKG = ROOT / "docs" / "教师上课资料包_信息系统复习课"
GUIDE_MD = PKG / "学生导学案.md"
GUIDE_DOCX = PKG / "学生导学案.docx"
GUIDE_DOCX_FALLBACK = PKG / "学生导学案_新版.docx"


def set_font(run, size=11, bold=False):
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.bold = bold


def clean_inline(text: str) -> str:
    return text.replace("`", "")


def add_markdown_docx(md_path: Path, out_path: Path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    for name, size in [("Title", 18), ("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        styles[name].font.name = "Microsoft YaHei"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(line[2:].strip()), 18, True)
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_paragraph(style="Heading 1")
            set_font(p.add_run(line[3:].strip()), 14, True)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            set_font(p.add_run(line[4:].strip()), 12, True)
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                stripped = lines[i].strip()
                if not re.match(r"^\|\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", stripped):
                    table_lines.append(stripped)
                i += 1

            rows = [
                [clean_inline(cell.strip()) for cell in table_line.strip("|").split("|")]
                for table_line in table_lines
            ]
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, row in enumerate(rows):
                    for c_idx, text in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        set_font(cell.paragraphs[0].add_run(text), 10 if r_idx == 0 else 9.2, r_idx == 0)
                doc.add_paragraph()
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            set_font(p.add_run(clean_inline(line[2:].strip())), 10.5)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            set_font(p.add_run(clean_inline(re.sub(r"^\d+\.\s+", "", line))), 10.5)
            i += 1
            continue

        p = doc.add_paragraph()
        set_font(p.add_run(clean_inline(line)), 10.5)
        i += 1

    doc.save(out_path)


if __name__ == "__main__":
    try:
        add_markdown_docx(GUIDE_MD, GUIDE_DOCX)
        generated_docx = GUIDE_DOCX
    except PermissionError:
        generated_docx = GUIDE_DOCX_FALLBACK
        add_markdown_docx(GUIDE_MD, generated_docx)
    shutil.copy2(GUIDE_MD, ROOT / "docs" / "信息系统复习课导学案.md")
    shutil.copy2(generated_docx, ROOT / "docs" / "信息系统复习课导学案.docx")
    shutil.copy2(PKG / "老师上课带学生用_活动页面.html", ROOT / "docs" / "信息系统复习课活动页面.html")
    print("generated", generated_docx)
